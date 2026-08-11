"""BOSS resume capability runner: atomic modules, no fixed workflow."""
from __future__ import annotations
import argparse, hashlib, json, os, time, zipfile
from pathlib import Path
DEFAULT_MESSAGE="方便发送一份简历给我吗？"
INPUT_AIDS={"bosschat-global-input","boss-chat-editor-input"}

def fail(code,message,**extra): raise RuntimeError(json.dumps({"code":code,"message":message,**extra},ensure_ascii=False))
def norm(s): return str(s or "").replace("\r\n","\n").replace("\r","\n").strip()
def deps():
 if os.name!='nt': fail('WINDOWS_REQUIRED','Windows required')
 try:
  from pywinauto import Desktop, keyboard, mouse
 except Exception as e: fail('RUNTIME_MISSING',str(e))
 return Desktop,keyboard,mouse
def window():
 Desktop,_,_=deps(); ws=Desktop(backend='uia').windows(title='BOSS直聘',visible_only=True)
 if len(ws)!=1: fail('BOSS_WINDOW_NOT_UNIQUE',f'count={len(ws)}')
 w=ws[0]
 if len(w.descendants())<100: fail('ACCESSIBILITY_UNAVAILABLE','semantic UI tree incomplete')
 return w
def invoke(c):
 for fn in (lambda:c.iface_invoke.Invoke(),lambda:c.invoke(),lambda:c.iface_legacyIAccessible.DoDefaultAction(),lambda:c.click_input()):
  try: fn(); return
  except Exception: pass
 fail('SEMANTIC_ACTION_FAILED','control could not be invoked')
def text_nodes(c): return [norm(x.window_text()) for x in c.descendants(control_type='Text') if norm(x.window_text())]
def editor(w):
 xs=[c for c in w.descendants(control_type='Group') if c.is_visible() and c.element_info.automation_id in INPUT_AIDS]
 if len(xs)!=1: fail('EDITOR_NOT_UNIQUE',f'count={len(xs)}')
 return xs[0]
def open_message_and_job(job):
 w=window(); nav=[c for c in w.descendants(control_type='Hyperlink') if c.is_visible() and norm(c.window_text()).startswith('消息')]
 if len(nav)==1: invoke(nav[0]);time.sleep(.6);w=window()
 elif len(nav)>1: fail('MESSAGE_NAV_NOT_UNIQUE',f'count={len(nav)}')
 top=w.rectangle().top;anchors=[c for c in w.descendants(control_type='Text') if c.is_visible() and c.rectangle().top<top+180 and (norm(c.window_text())=='全部职位' or ' _ ' in norm(c.window_text()))]
 if len(anchors)!=1: fail('JOB_ANCHOR_NOT_UNIQUE',f'count={len(anchors)}')
 if norm(anchors[0].window_text())==norm(job): return w
 invoke(anchors[0]);time.sleep(.3);w=window();opts=[]
 for c in w.descendants(control_type='Text'):
  try:
   t=norm(c.window_text())
   if c.is_visible() and c.parent().element_info.control_type=='ListItem' and (t=='全部职位' or ' _ ' in t):opts.append(c)
  except:pass
 exact=[c for c in opts if norm(c.window_text())==norm(job)]
 if len(exact)!=1:fail('JOB_NOT_EXACT',f'count={len(exact)}',candidates=[norm(x.window_text()) for x in opts])
 invoke(exact[0]);time.sleep(.7);return window()
def document_text_control(w):
 docs=[]
 for d in w.descendants(control_type='Document'):
  try:
   text=d.iface_text.DocumentRange.GetText(-1)
   if text: docs.append((len(text),d))
  except Exception: pass
 if not docs: fail('MESSAGE_TEXT_DOCUMENT_NOT_FOUND','BOSS UIA Document TextPattern is unavailable')
 return max(docs,key=lambda x:x[0])[1]
def job_title(job): return norm(job).split(' _ ',1)[0]
def candidate_phrase(candidate,job): return norm(candidate)+' '+job_title(job)
def find_text_ranges(document,needle):
 root=document.iface_text.DocumentRange;ranges=[];cursor=root.Clone()
 while True:
  found=cursor.FindText(needle,False,True)
  if not found: break
  ranges.append((document,found));cursor.MoveEndpointByRange(0,found,1)
  if cursor.CompareEndpoints(0,root,1)>=0: break
 return ranges
def candidate_text_ranges(w,candidate,job):
 d=document_text_control(w);composite=find_text_ranges(d,candidate_phrase(candidate,job))
 if composite:return composite
 # BOSS can show a conversation-specific title such as “兼职·AI员工培训师…”
 # while the top filter is “AI员工培训师…”. If the composite differs, search
 # the exact name and retain only ranges enclosed by a message ListItem; this
 # excludes chat header/history occurrences without opening any conversation.
 rows=[];seen=set()
 for pair in find_text_ranges(d,norm(candidate)):
  try:
   item=enclosing_list_item(pair[1]);r=item.rectangle()
   # Conversation rows occupy the left message-list pane. Chat-message bubbles
   # are in the right pane and must never participate in candidate selection.
   if not (350<=r.left and r.right<=1250 and r.width()>500 and r.height()>100):continue
   key=(r.left,r.top,r.right,r.bottom)
   if key not in seen:seen.add(key);rows.append(pair)
  except RuntimeError:pass
 return rows
def enclosing_list_item(text_range):
 try:
  from pywinauto.uia_element_info import UIAElementInfo
  from pywinauto.controls.uiawrapper import UIAWrapper
  item=UIAWrapper(UIAElementInfo(text_range.GetEnclosingElement()))
  while item.element_info.control_type!='ListItem': item=item.parent()
  return item
 except Exception as exc: fail('MESSAGE_ROW_FROM_TEXT_FAILED',str(exc))
def document_text(w):
 d=document_text_control(w);return d,d.iface_text.DocumentRange.GetText(-1)
def locate_candidate_with_loading(w,candidate,job,max_load_rounds=12,max_seconds=45):
 # Search the currently exposed full TextPattern first. Only a miss authorizes
 # semantic pagination through the unique “滚动加载更多” TextRange.
 started=time.time();rounds=0;seen=set()
 while True:
  matches=candidate_text_ranges(w,candidate,job)
  if len(matches)==1:return matches,rounds
  if len(matches)>1:fail('CANDIDATE_NOT_UNIQUE_AFTER_JOB_FILTER',f'count={len(matches)}',candidate=candidate,job=job,load_rounds=rounds)
  d,text=document_text(w);fingerprint=hashlib.sha256(text.encode('utf-8','ignore')).hexdigest()
  if fingerprint in seen:fail('MESSAGE_LIST_NO_PROGRESS','UIA message-list text repeated without locating candidate',candidate=candidate,job=job,load_rounds=rounds)
  seen.add(fingerprint)
  if '没有更多了' in text:fail('CANDIDATE_NOT_FOUND_AFTER_UIA_PAGINATION','candidate absent and BOSS reports no more conversations',candidate=candidate,job=job,load_rounds=rounds)
  if rounds>=max_load_rounds or time.time()-started>=max_seconds:fail('CANDIDATE_SEARCH_LIMIT_REACHED','bounded UIA pagination limit reached',candidate=candidate,job=job,load_rounds=rounds)
  more=find_text_ranges(d,'滚动加载更多')
  if len(more)==0:fail('CANDIDATE_NOT_FOUND_IN_UIA_LIST','candidate absent and no UIA load-more range is available',candidate=candidate,job=job,load_rounds=rounds)
  if len(more)>1:fail('MESSAGE_LOAD_MORE_NOT_UNIQUE',f'count={len(more)}',load_rounds=rounds)
  try:more[0][1].ScrollIntoView(True)
  except Exception as exc:fail('MESSAGE_PAGINATION_FAILED',str(exc),load_rounds=rounds)
  rounds+=1;deadline=min(started+max_seconds,time.time()+5);changed=False
  while time.time()<deadline:
   time.sleep(.4);w=window();_,after=document_text(w)
   after_fp=hashlib.sha256(after.encode('utf-8','ignore')).hexdigest()
   if after_fp!=fingerprint:changed=True;break
  if not changed:fail('MESSAGE_LIST_NO_PROGRESS','UIA load-more did not change the message-list text',candidate=candidate,job=job,load_rounds=rounds)
def open_conversation(job,candidate):
 w=open_message_and_job(job);matches,load_rounds=locate_candidate_with_loading(w,candidate,job)
 _,rng=matches[0]
 # TextPattern exposes virtualized/off-screen rows. Scroll the exact semantic
 # range into view, then reacquire it because Chromium can recycle ListItems.
 try: rng.ScrollIntoView(True)
 except Exception as exc: fail('MESSAGE_ROW_SCROLL_FAILED',str(exc))
 time.sleep(.55);w=window();matches=candidate_text_ranges(w,candidate,job)
 if len(matches)!=1: fail('CANDIDATE_NOT_UNIQUE_AFTER_SCROLL',f'count={len(matches)}',candidate=candidate,job=job,load_rounds=load_rounds)
 item=enclosing_list_item(matches[0][1])
 if not item.is_visible(): fail('MESSAGE_ROW_NOT_VISIBLE','TextPattern match did not materialize a visible ListItem')
 item.click_input();time.sleep(.8);w=window()
 try: editor(w)
 except RuntimeError: fail('CONVERSATION_NOT_READY','direct UIA list selection did not open a unique chat editor',candidate=candidate,job=job)
 return w
def messages(w):
 out=[]
 for x in w.descendants(control_type='ListItem'):
  aid=str(x.element_info.automation_id)
  if aid.startswith('mid-'):out.append({'item':x,'id':aid,'texts':text_nodes(x),'visible':x.is_visible()})
 return out
def state(w):
 ms=messages(w)
 attachments=[{'message_id':x['id'],'visible':x['visible'],'texts':x['texts']} for x in ms if '点击预览附件简历' in x['texts']]
 attachment_indexes=[i for i,x in enumerate(ms) if '点击预览附件简历' in x['texts']]
 pending=[]
 for i,x in enumerate(ms):
  if '拒绝' in x['texts'] and '同意' in x['texts']:
   # BOSS may leave the old agree card visible after acceptance. A later attachment
   # message resolves that request and prevents a second click.
   if not any(j>i for j in attachment_indexes):pending.append({'message_id':x['id'],'visible':x['visible'],'texts':x['texts']})
 return {'platform_requested':any('简历请求已发送' in x['texts'] for x in ms),'pending_attachment_requests':pending,'attachments':attachments}
def runtime_root():
 p=Path(os.environ.get('LOCALAPPDATA') or Path.home())/'CyberNuwa'/'boss-resume-request-collection';p.mkdir(parents=True,exist_ok=True);return p
def ledger_key(job,candidate,mode,request_id): return hashlib.sha256(json.dumps([job,candidate,mode,request_id],ensure_ascii=False,separators=(',',':')).encode()).hexdigest()
def ledger():
 p=runtime_root()/'actions.jsonl'; keys=set()
 if p.exists():
  for line in p.read_text(encoding='utf-8').splitlines():
   try:keys.add(json.loads(line)['key'])
   except:pass
 return p,keys
def append_ledger(path,key,result):
 with path.open('a',encoding='utf-8') as f:f.write(json.dumps({'key':key,'at':time.time(),'result':result},ensure_ascii=False)+'\n');f.flush();os.fsync(f.fileno())
def receipt(job,candidate,mode,status,quota=False,evidence=None): return {'job_key':job,'candidate_key':candidate,'candidate_name':candidate,'request_mode':mode,'request_status':status,'quota_consumed':quota,'resume_status':'NOT_RECEIVED','original_file':'','file_size':0,'sha256':'','parse_status':'NOT_STARTED','evidence':evidence or []}
def fixed_text_infos(w,title):
 # Native UIA property condition, matching the job-publishing runtime. This is
 # a fixed selector lookup, not a client-side visible-control name scan.
 infos=w.element_info.descendants(cache_enable=False,title=title,control_type='Text')
 from pywinauto.controls.uiawrapper import UIAWrapper
 return [UIAWrapper(info) for info in infos]
def request_resume_control(w):
 # The request action is the first semantic Text in the fixed action strip
 # immediately above the unique editor. Resolve by editor scope + geometry,
 # then validate the expected role; never search all visible Texts for 求简历.
 ed=editor(w);er=ed.rectangle();band=[]
 for c in w.descendants(control_type='Text'):
  try:
   r=c.rectangle()
   if c.is_visible() and er.left<=r.left<er.right and 0<=er.top-r.bottom<=80:band.append(c)
  except Exception:pass
 band.sort(key=lambda c:c.rectangle().left)
 if len(band)<3:fail('REQUEST_ACTION_STRIP_INVALID','editor action strip is incomplete',count=len(band))
 target=band[0]
 if norm(target.window_text())!='求简历':fail('REQUEST_ACTION_ROLE_MISMATCH','first editor action is not 求简历',actual=norm(target.window_text()))
 return target
def pending_accept_control(item):
 # A pending attachment ListItem has one file Image and exactly two action Text
 # siblings. The right-hand action is fixed as Accept. Select by structure/order,
 # then read back its role; never search descendants by name 同意.
 descendants=item.descendants();images=[c for c in descendants if c.element_info.control_type=='Image'];actions=[c for c in descendants if c.element_info.control_type=='Text']
 actions.sort(key=lambda c:c.rectangle().left)
 if len(images)!=1 or len(actions)!=2:fail('PENDING_ATTACHMENT_STRUCTURE_INVALID','expected one Image and two action Text siblings',image_count=len(images),action_count=len(actions))
 if [norm(c.window_text()) for c in actions]!=['拒绝','同意']:fail('PENDING_ATTACHMENT_ROLE_MISMATCH','pending action order changed',actual=[norm(c.window_text()) for c in actions])
 return actions[1]
def confirmation_controls(w):
 # Anchor the one authorized prompt with a native fixed selector. Its two action
 # Texts occupy the same modal band; resolve left/right by structure and validate.
 prompts=fixed_text_infos(w,'确定向牛人索取简历吗？')
 prompts=[c for c in prompts if c.is_visible()]
 if len(prompts)!=1:fail('CONFIRMATION_DIALOG_NOT_EXACT','expected one 确定向牛人索取简历吗？ prompt; do not retry',count=len(prompts))
 pr=prompts[0].rectangle();actions=[]
 for c in w.descendants(control_type='Text'):
  try:
   r=c.rectangle()
   if c.is_visible() and r.top>pr.bottom and 20<=r.top-pr.bottom<=90 and r.left>=pr.left and r.right<=pr.right+100:actions.append(c)
  except Exception:pass
 actions.sort(key=lambda c:c.rectangle().left)
 role=[norm(c.window_text()) for c in actions]
 pairs=[actions[i:i+2] for i in range(max(0,len(actions)-1)) if [norm(x.window_text()) for x in actions[i:i+2]]==['取消','确定']]
 if len(pairs)!=1:fail('CONFIRMATION_CONTROLS_NOT_EXACT','modal action pair is not structurally unique; do not retry',roles=role)
 return prompts[0],pairs[0][0],pairs[0][1]
def request_platform(job,candidate,request_id):
 w=open_conversation(job,candidate);s=state(w);p,keys=ledger();key=ledger_key(job,candidate,'PLATFORM',request_id)
 if key in keys or s['platform_requested'] or s['attachments']: return receipt(job,candidate,'PLATFORM','ALREADY_HANDLED',False,[s])
 target=request_resume_control(w)
 before={x['id'] for x in messages(w)};invoke(target)
 # Verified 2026-08-05 on BOSS desktop: clicking 求简历 opens the semantic
 # confirmation “确定向牛人索取简历吗？”, with unique 取消/确定 Text controls.
 # PLATFORM mode explicitly authorizes this one confirmation, but no other dialog.
 deadline=time.time()+3;controls=None
 while time.time()<deadline:
  time.sleep(.2);w=window()
  try:controls=confirmation_controls(w);break
  except RuntimeError:pass
 if controls is None:fail('CONFIRMATION_DIALOG_NOT_EXACT','authorized request confirmation did not become structurally ready; do not retry')
 _,_,confirm=controls;invoke(confirm);time.sleep(1.2);w=window()
 remaining=[c for c in fixed_text_infos(w,'确定向牛人索取简历吗？') if c.is_visible()]
 if remaining:fail('COMMIT_UNKNOWN','confirmation remained after click; do not retry')
 after=messages(w);new=[x for x in after if x['id'] not in before]
 request_marks=[x for x in new if '简历请求已发送' in x['texts']]
 sent_prompts=[x for x in new if any('方便发一份简历过来吗' in t for t in x['texts'])]
 if len(request_marks)!=1: fail('COMMIT_UNKNOWN','request outcome not uniquely verified; do not retry',new_messages=[{'id':x['id'],'texts':x['texts']} for x in new])
 out=receipt(job,candidate,'PLATFORM','REQUESTED_VERIFIED',True,[{'message_id':request_marks[0]['id']},{'platform_prompt_message_ids':[x['id'] for x in sent_prompts]}]);append_ledger(p,key,out);return out
def request_message(job,candidate,request_id,message):
 w=open_conversation(job,candidate);p,keys=ledger();key=ledger_key(job,candidate,'MESSAGE',request_id)
 if key in keys or any(norm(message) in x['texts'] for x in messages(w)):return receipt(job,candidate,'MESSAGE','ALREADY_HANDLED',False)
 _,keyboard,_=deps();g=editor(w);w.set_focus();g.set_focus();keyboard.send_keys('^a{BACKSPACE}');keyboard.send_keys(message,with_spaces=True);actual=norm(g.legacy_properties().get('Value',''))
 if actual!=norm(message):fail('DRAFT_MISMATCH','draft readback mismatch')
 before={x['id'] for x in messages(w)};keyboard.send_keys('{ENTER}');time.sleep(.7);new=[x for x in messages(window()) if x['id'] not in before and norm(message) in x['texts']]
 if len(new)!=1:fail('COMMIT_UNKNOWN','send outcome not uniquely verified; do not retry')
 out=receipt(job,candidate,'MESSAGE','SENT_VERIFIED',False,[{'message_id':new[0]['id']}]);append_ledger(p,key,out);return out
def inspect_state(job,candidate):return state(open_conversation(job,candidate))
def accept_pending_attachment(job,candidate,request_message_id=None):
 w=open_conversation(job,candidate);before_state=state(w);pending=before_state['pending_attachment_requests']
 if request_message_id:pending=[x for x in pending if x['message_id']==request_message_id]
 if not pending:
  return {'accept_status':'ALREADY_ACCEPTED' if before_state['attachments'] else 'NO_PENDING_REQUEST','attachments':before_state['attachments'],'evidence':[]}
 if len(pending)!=1:fail('PENDING_ATTACHMENT_SELECTION_REQUIRED','multiple pending attachment requests',pending=pending)
 aid=pending[0]['message_id'];items=[x for x in w.descendants(control_type='ListItem') if str(x.element_info.automation_id)==aid]
 if len(items)!=1:fail('PENDING_ATTACHMENT_NOT_EXACT',f'count={len(items)}')
 target=pending_accept_control(items[0])
 before={x['id'] for x in messages(w)};invoke(target);time.sleep(1.0);after=messages(window())
 new_attachments=[x for x in after if x['id'] not in before and '点击预览附件简历' in x['texts']]
 if len(new_attachments)!=1:fail('COMMIT_UNKNOWN','accept outcome not uniquely verified; do not retry',new_messages=[{'id':x['id'],'texts':x['texts']} for x in after if x['id'] not in before])
 return {'accept_status':'ACCEPTED_VERIFIED','pending_message_id':aid,'attachment_message_id':new_attachments[0]['id'],'evidence':[{'message_id':new_attachments[0]['id'],'texts':new_attachments[0]['texts']}]}
def attachment_download_point(item):
 # BOSS 1.7.4 does not expose the download glyph as a UIA descendant. Resolve
 # the verified attachment card by its semantic ListItem, then use the card's
 # fixed internal structure: one preview Text, one timestamp Text sibling and
 # one file Image. The glyph is at a DPI-scaled offset from the timestamp, so
 # moving or resizing the BOSS window does not change this relationship.
 descendants=item.descendants();previews=[c for c in descendants if c.element_info.control_type=='Text' and norm(c.window_text())=='点击预览附件简历']
 if len(previews)!=1:fail('ATTACHMENT_STRUCTURE_INVALID','preview anchor is not unique',count=len(previews))
 texts=[c for c in descendants if c.element_info.control_type=='Text' and c is not previews[0]]
 images=[c for c in descendants if c.element_info.control_type=='Image']
 if len(texts)!=1 or len(images)!=1:fail('ATTACHMENT_STRUCTURE_INVALID','timestamp/image sibling structure is not exact',text_count=len(texts),image_count=len(images))
 timestamp=texts[0];tr=timestamp.rectangle();pr=previews[0].rectangle();ir=images[0].rectangle()
 unit=max(1,tr.height())
 # Structural sanity checks only; no whole-bubble ratio and no window geometry.
 if not (ir.left<pr.left<tr.left and tr.top<ir.top<pr.top and 2.0*unit<=ir.width()<=3.5*unit and 2.0*unit<=ir.height()<=3.5*unit):fail('ATTACHMENT_STRUCTURE_INVALID','attachment anchors have unexpected ordering')
 return round(pr.right+7.37*unit),round(tr.bottom+3.22*unit)
def download_received(job,candidate,output_dir,attachment_message_id=None):
 w=open_conversation(job,candidate);ats=state(w)['attachments']
 if attachment_message_id:ats=[x for x in ats if x['message_id']==attachment_message_id]
 if not ats:return {'resume_status':'NOT_RECEIVED','evidence':[]}
 if len(ats)!=1:fail('ATTACHMENT_SELECTION_REQUIRED','multiple attachments',attachments=ats)
 aid=ats[0]['message_id'];items=[x for x in w.descendants(control_type='ListItem') if str(x.element_info.automation_id)==aid]
 if len(items)!=1:fail('ATTACHMENT_BUBBLE_NOT_EXACT',f'count={len(items)}')
 item=items[0]
 try:item.iface_scroll_item.ScrollIntoView();time.sleep(.4)
 except:pass
 out=Path(output_dir);out.mkdir(parents=True,exist_ok=True);before={p.resolve() for p in out.iterdir() if p.is_file()}
 # Fixed structural entry: never scan for a visible control named 下载 / 下载简历 / 下载附件.
 # Reacquire the card after scrolling because Chromium may recycle wrappers.
 try:
  _,_,mouse=deps();fresh=[]
  for _ in range(8):
   w=window();fresh=[x for x in w.descendants(control_type='ListItem') if str(x.element_info.automation_id)==aid]
   if len(fresh)==1:break
   time.sleep(.2)
  if len(fresh)!=1:fail('ATTACHMENT_BUBBLE_NOT_EXACT',f'count={len(fresh)}')
  point=attachment_download_point(fresh[0]);mouse.click(coords=point)
 except Exception as e:fail('DOWNLOAD_CONTROL_UNAVAILABLE',str(e))
 time.sleep(.5)
 # Native save dialog: navigate its address bar to output_dir first, while
 # preserving the platform filename. Writing a full path into the filename box
 # is unreliable and can silently save to Desktop, so it is explicitly banned.
 Desktop,keyboard,_=deps();deadline=time.time()+4;dialogs=[]
 while time.time()<deadline:
  dialogs=Desktop(backend='uia').windows(title='下载',visible_only=True)
  if dialogs:break
  time.sleep(.2)
 if len(dialogs)!=1:fail('SAVE_DIALOG_NOT_EXACT',f'download save dialog count={len(dialogs)}')
 d=dialogs[0];d.set_focus();filename_edits=[c for c in d.descendants(control_type='Edit') if str(c.element_info.automation_id)=='1001']
 if len(filename_edits)!=1:fail('SAVE_DIALOG_INVALID','filename input is not unique')
 original=norm(filename_edits[0].get_value()) or 'resume.pdf'
 if Path(original).name!=original:fail('SAVE_DIALOG_INVALID','platform filename unexpectedly contains a path')
 keyboard.send_keys('%d');time.sleep(.15);keyboard.send_keys(str(out.resolve()),with_spaces=True);keyboard.send_keys('{ENTER}');time.sleep(.7)
 dialogs=Desktop(backend='uia').windows(title='下载',visible_only=True)
 if len(dialogs)!=1:fail('SAVE_DIALOG_NOT_EXACT',f'download save dialog count={len(dialogs)}')
 d=dialogs[0]
 address=[c for c in d.descendants(control_type='ToolBar') if norm(c.window_text()).startswith('地址:')]
 expected=os.path.normcase(os.path.normpath(str(out.resolve())))
 if len(address)!=1 or expected not in os.path.normcase(os.path.normpath(norm(address[0].window_text()).removeprefix('地址:').strip())):fail('SAVE_DIRECTORY_NOT_CONFIRMED','save dialog did not enter output_dir')
 filename_edits=[c for c in d.descendants(control_type='Edit') if str(c.element_info.automation_id)=='1001']
 if len(filename_edits)!=1 or norm(filename_edits[0].get_value())!=original:fail('SAVE_FILENAME_CHANGED','platform filename changed while navigating')
 saves=[c for c in d.descendants(control_type='Button') if str(c.element_info.automation_id)=='1' and norm(c.window_text()).startswith('保存')]
 if len(saves)!=1:fail('SAVE_DIALOG_INVALID','save button is not unique')
 invoke(saves[0]);time.sleep(.8)
 deadline=time.time()+20
 while time.time()<deadline:
  created=[p for p in out.iterdir() if p.is_file() and p.resolve() not in before and p.suffix.lower() in {'.pdf','.docx'} and not p.name.endswith(('.crdownload','.tmp'))]
  if len(created)==1:return validate_file(created[0])
  if len(created)>1:fail('DOWNLOAD_NOT_UNIQUE','multiple files created',files=[str(x) for x in created])
  time.sleep(.5)
 fail('DOWNLOAD_NOT_OBSERVED','no original PDF/DOCX appeared in output_dir')
def validate_file(path):
 p=Path(path);data=p.read_bytes() if p.is_file() else b''
 if not data:fail('FILE_EMPTY_OR_MISSING',str(p))
 if data.startswith(b'%PDF-'):real='PDF'
 elif zipfile.is_zipfile(p):
  with zipfile.ZipFile(p) as z:n=set(z.namelist())
  real='DOCX' if {'[Content_Types].xml','word/document.xml'}<=n else 'ZIP_OTHER'
 else:real='UNKNOWN'
 expected={'.pdf':'PDF','.docx':'DOCX'}.get(p.suffix.lower())
 if real!=expected:fail('FILE_TYPE_MISMATCH','extension and signature differ',expected=expected,actual=real)
 h=hashlib.sha256(data).hexdigest();return {'resume_status':'DOWNLOADED','original_file':str(p.resolve()),'file_size':len(data),'sha256':h,'parse_status':'NOT_STARTED','evidence':[{'format':real}]}
def parse_file(path):
 p=Path(path);v=validate_file(p)
 if p.suffix.lower()=='.pdf':
  try:
   from pypdf import PdfReader;text='\n'.join((x.extract_text() or '') for x in PdfReader(str(p)).pages)
  except Exception as e:fail('PARSE_FAILED',str(e))
 else:
  try:
   from docx import Document;text='\n'.join(x.text for x in Document(str(p)).paragraphs)
  except Exception as e:fail('PARSE_FAILED',str(e))
 v.update({'resume_status':'PARSED','parse_status':'PARSED','text_length':len(text)});return v
def main():
 p=argparse.ArgumentParser();sp=p.add_subparsers(dest='cmd',required=True);sp.add_parser('runtime')
 for name in ['inspect-state','request-platform','request-message','accept-pending','download-received','collect']:
  x=sp.add_parser(name);x.add_argument('--job',required=True);x.add_argument('--candidate',required=True)
  if name in ['request-platform','request-message']:x.add_argument('--request-id',required=True)
  if name=='request-message':x.add_argument('--message-file')
  if name=='accept-pending':x.add_argument('--request-message-id')
  if name in ['download-received','collect']:x.add_argument('--output-dir',required=True);x.add_argument('--attachment-message-id')
 for name in ['validate-file','parse-file']:
  x=sp.add_parser(name);x.add_argument('--file',required=True)
 a=p.parse_args()
 try:
  if a.cmd=='runtime':res={'python':os.sys.executable,'pywinauto':__import__('pywinauto').__version__}
  elif a.cmd=='inspect-state':res=inspect_state(a.job,a.candidate)
  elif a.cmd=='request-platform':res=request_platform(a.job,a.candidate,a.request_id)
  elif a.cmd=='request-message':res=request_message(a.job,a.candidate,a.request_id,Path(a.message_file).read_text(encoding='utf-8') if a.message_file else DEFAULT_MESSAGE)
  elif a.cmd=='accept-pending':res=accept_pending_attachment(a.job,a.candidate,a.request_message_id)
  elif a.cmd=='download-received':res=download_received(a.job,a.candidate,a.output_dir,a.attachment_message_id)
  elif a.cmd=='validate-file':res=validate_file(a.file)
  elif a.cmd=='parse-file':res=parse_file(a.file)
  else:
   d=download_received(a.job,a.candidate,a.output_dir,a.attachment_message_id);res=parse_file(d['original_file']) if d.get('original_file') else d
  print(json.dumps({'ok':True,'result':res},ensure_ascii=False));return 0
 except Exception as e:
  try:err=json.loads(str(e))
  except:err={'code':type(e).__name__,'message':str(e)}
  print(json.dumps({'ok':False,'error':err},ensure_ascii=False),file=os.sys.stderr);return 1
if __name__=='__main__':raise SystemExit(main())
